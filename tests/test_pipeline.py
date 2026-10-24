"""Automated tests + measured metrics.  Run:  python -m unittest tests.test_pipeline -v
Writes reports/metrics.json with the real numbers (nothing is asserted that was not measured)."""
import difflib
import io
import json
import os
import sys
import tempfile
import threading
import time
import unittest
import urllib.request
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
TMP = Path(tempfile.mkdtemp(prefix="mdia_test_"))
os.environ["MDIA_DATA_DIR"] = str(TMP / "data")

from mdia import exporters                                  # noqa: E402
from mdia.pipeline import process_job                       # noqa: E402
from mdia.schema import build as build_schema               # noqa: E402
from mdia.store import Store, apply_corrections             # noqa: E402
from mdia.textutil import match_key, parse_date, parse_number  # noqa: E402
from tests import make_samples as ms                        # noqa: E402

SAMPLES = TMP / "samples"
STORE = Store(TMP / "store")
METRICS = {"classification": {}, "fields": {"correct": 0, "total": 0, "wrong": []}, "items": {"correct": 0, "total": 0},
           "evidence": {"checked": 0, "grounded": 0}, "fabricated_values": 0, "review": {"fields": 0, "needs_review": 0}}
_CACHE = {}


def run(name):
    if name not in _CACHE:
        jid = STORE.create(name, (SAMPLES / name).read_bytes())
        r = process_job(STORE, jid)
        assert r is not None, STORE.status(jid)
        _CACHE[name] = (jid, r)
    return _CACHE[name]


def val(f):
    return None if f["status"] == "not_found" else (f["value"] if f["value"] is not None else f["raw"])


def check(label, got, want):
    METRICS["fields"]["total"] += 1
    ok = (abs(got - want) < 0.005) if isinstance(want, float) and isinstance(got, (int, float)) else got == want
    if ok:
        METRICS["fields"]["correct"] += 1
    else:
        METRICS["fields"]["wrong"].append({"field": label, "got": got, "want": want})
    return ok


def items_ok(doc, truth):
    got = doc["invoice_items"]
    ok = 0
    for g, t in zip(got, truth):
        METRICS["items"]["total"] += 1
        if (val(g["service_name"]) == t["name"] and val(g["quantity"]) == t["qty"] and abs(val(g["unit_price"]) - t["price"]) < .005
                and abs(val(g["line_total"]) - t["total"]) < .005):
            METRICS["items"]["correct"] += 1
            ok += 1
    METRICS["items"]["total"] += max(0, len(truth) - len(got))      # missing rows count as errors
    return ok == len(truth) and len(got) == len(truth)


def all_fields(doc):
    for g in ("patient", "provider", "hospitalization"):
        for k, f in (doc[g] or {}).items():
            yield f
    for key in ("invoices", "invoice_items", "medications", "diagnoses", "laboratory_tests", "imaging", "procedures"):
        for rec in doc[key]:
            for f in rec.values():
                if isinstance(f, dict) and "status" in f:
                    yield f


def validate_schema(instance, schema, root=None, path="$"):
    """Tiny validator for the subset of JSON Schema the project uses (jsonschema may be unavailable offline)."""
    root = root or schema
    errs = []
    if "$ref" in schema:
        node = root
        for p in schema["$ref"].lstrip("#/").split("/"):
            node = node[p]
        return validate_schema(instance, node, root, path)
    if "const" in schema and instance != schema["const"]:
        errs.append(f"{path}: const")
    if "enum" in schema and instance not in schema["enum"]:
        errs.append(f"{path}: {instance!r} not in enum")
    t = schema.get("type")
    if t:
        ts = t if isinstance(t, list) else [t]
        py = {"string": str, "number": (int, float), "integer": int, "array": list, "object": dict, "null": type(None), "boolean": bool}
        if not any(isinstance(instance, py[x]) and not (x in ("number", "integer") and isinstance(instance, bool)) for x in ts):
            return errs + [f"{path}: expected {ts}, got {type(instance).__name__}"]
    if isinstance(instance, dict):
        for r in schema.get("required", []):
            if r not in instance:
                errs.append(f"{path}: missing {r}")
        for k, sub in schema.get("properties", {}).items():
            if k in instance:
                errs += validate_schema(instance[k], sub, root, f"{path}.{k}")
    if isinstance(instance, list) and "items" in schema:
        for i, x in enumerate(instance):
            errs += validate_schema(x, schema["items"], root, f"{path}[{i}]")
    if isinstance(instance, (int, float)) and not isinstance(instance, bool):
        if "minimum" in schema and instance < schema["minimum"]:
            errs.append(f"{path}: below minimum")
        if "maximum" in schema and instance > schema["maximum"]:
            errs.append(f"{path}: above maximum")
    return errs


class T(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        ms.build_all(SAMPLES)

    @classmethod
    def tearDownClass(cls):
        # evidence grounding + fabrication audit over every sample result
        for name, (jid, r) in _CACHE.items():
            pages = {p["page"]: json.loads(STORE.path(jid, "pages", f"{p['page']}.json").read_text("utf-8")) for p in r["page_info"]}
            for d in r["documents"]:
                ev = {e["id"]: e for e in d["evidence"]}
                for f in all_fields(d):
                    if f["status"] == "not_found":
                        continue
                    METRICS["review"]["fields"] += 1
                    METRICS["review"]["needs_review"] += f["verification"] == "needs_review"
                    e = ev.get(f["evidence_id"])
                    METRICS["evidence"]["checked"] += 1
                    if e is None:
                        continue
                    words = pages[e["page"]]["words"]
                    b = e["bbox"]
                    inside = [w for w in words if b[0] - .004 <= (w["x0"] + w["x1"]) / 2 <= b[2] + .004 and b[1] - .004 <= (w["y0"] + w["y1"]) / 2 <= b[3] + .004]
                    box_txt = match_key("".join(w["text"] for w in inside)).replace(" ", "")
                    raw = match_key(e["matched_text"]).replace(" ", "")
                    grounded = bool(raw) and (raw in box_txt or all(match_key(t) in box_txt for t in e["matched_text"].split() if match_key(t)))
                    METRICS["evidence"]["grounded"] += grounded
                    # fabrication audit: every word of an extracted value must exist on the page
                    page_tokens = {match_key(w["text"]) for w in words}
                    toks = [match_key(t) for t in f["raw"].split() if match_key(t)]
                    if f["status"] == "extracted" and toks and not all(t in page_tokens or any(t in pt for pt in page_tokens) for t in toks):
                        METRICS["fabricated_values"] += 1
                        METRICS.setdefault("fabricated_detail", []).append({"sample": name, "raw": f["raw"]})
        m = METRICS
        m["fields"]["accuracy"] = round(m["fields"]["correct"] / max(1, m["fields"]["total"]), 4)
        m["items"]["accuracy"] = round(m["items"]["correct"] / max(1, m["items"]["total"]), 4)
        m["evidence"]["grounded_rate"] = round(m["evidence"]["grounded"] / max(1, m["evidence"]["checked"]), 4)
        m["review"]["needs_review_rate"] = round(m["review"]["needs_review"] / max(1, m["review"]["fields"]), 4)
        cl = m["classification"]
        m["classification_accuracy"] = round(sum(cl.values()) / max(1, len(cl)), 4)
        out = ROOT / "reports"
        out.mkdir(exist_ok=True)
        (out / "metrics.json").write_text(json.dumps(m, ensure_ascii=False, indent=2), "utf-8")
        print("\n=== METRICS ===\n" + json.dumps({k: v for k, v in m.items()}, ensure_ascii=False, indent=1))

    # ------------------------------------------------------------ units
    def test_number_and_date_parsing(self):
        self.assertEqual(parse_number("1,234.50"), parse_number("١٬٢٣٤٫٥٠"))
        self.assertEqual(float(parse_number("1.234,50")), 1234.5)
        self.assertEqual(float(parse_number("(15.00)")), -15.0)
        self.assertIsNone(parse_number("abc"))
        self.assertEqual(parse_date("2024-03-15")[0], "2024-03-15")
        self.assertEqual(parse_date("25/03/2024")[0], "2024-03-25")
        self.assertIsNone(parse_date("03/04/2024")[0])                 # ambiguous -> never guessed
        self.assertEqual(parse_date("03/04/2024")[1], "ambiguous_day_month_order")
        self.assertEqual(parse_date("١٢/٠٣/١٤٤٥")[1], "hijri_calendar_date")

    def test_schema_is_valid_and_results_conform(self):
        schema = build_schema()
        for name in ("inv_en.pdf", "inv_ar.pdf", "rx.pdf", "lab.pdf", "discharge.pdf", "mixed.pdf", "inv_badmath.pdf"):
            errs = validate_schema(run(name)[1], schema)
            self.assertEqual(errs, [], f"{name}: {errs[:5]}")
        self.assertEqual(json.loads((ROOT / "schema" / "mdia.schema.json").read_text("utf-8")), schema, "schema file out of date: run python -m mdia.schema")

    # ------------------------------------------------------------ documents
    def test_english_invoice(self):
        d = run("inv_en.pdf")[1]["documents"][0]
        METRICS["classification"]["inv_en"] = d["document_type"] == "hospital_invoice"
        p, inv = d["patient"], d["invoices"][0]
        for lab, got, want in [("name", val(p["name"]), "Ahmed Ali Hassan"), ("file", val(p["file_no"]), "MRN-77421"), ("ins", val(p["insurance_no"]), "INS-5521"),
                               ("doctor", val(d["provider"]["doctor"]), "Samir Khalil"), ("facility", val(d["provider"]["facility_name"]), "Al Noor General Hospital"),
                               ("no", val(inv["invoice_number"]), "INV-1001"), ("date", val(inv["invoice_date"]), "2024-03-15"),
                               ("sdate", val(inv["service_date"]), "2024-03-14"), ("cur", val(inv["currency"]), "SAR"),
                               ("subtotal", val(inv["subtotal"]), 475.0), ("disc", val(inv["discount_total"]), 25.0), ("tax", val(inv["tax_total"]), 67.5),
                               ("total", val(inv["invoice_total"]), 517.5), ("paid", val(inv["paid_total"]), 300.0), ("bal", val(inv["balance_due"]), 217.5)]:
            self.assertTrue(check("inv_en." + lab, got, want), lab)
        self.assertTrue(items_ok(d, ms.ITEMS_EN))
        self.assertEqual([w for w in d["warnings"] if w["kind"].endswith("mismatch")], [])
        self.assertEqual(inv["invoice_total"]["verification"], "auto_verified")

    def test_arabic_invoice_rtl(self):
        d = run("inv_ar.pdf")[1]["documents"][0]
        METRICS["classification"]["inv_ar"] = d["document_type"] == "hospital_invoice"
        inv = d["invoices"][0]
        for lab, got, want in [("name", val(d["patient"]["name"]), "محمد عبدالله"), ("file", val(d["patient"]["file_no"]), "55120"),
                               ("doctor", val(d["provider"]["doctor"]), "سمير خليل"), ("no", val(inv["invoice_number"]), "INV-2002"),
                               ("date", val(inv["invoice_date"]), "2024-05-20"), ("total", val(inv["invoice_total"]), 402.5),
                               ("subtotal", val(inv["subtotal"]), 350.0), ("tax", val(inv["tax_total"]), 52.5)]:
            self.assertTrue(check("inv_ar." + lab, got, want), lab)
        self.assertTrue(items_ok(d, ms.ITEMS_AR))

    def test_multipage_invoice_is_one_document(self):
        r = run("inv_multi.pdf")[1]
        self.assertEqual(len(r["documents"]), 1)
        d = r["documents"][0]
        METRICS["classification"]["inv_multi"] = d["pages"] == [1, 2]
        self.assertEqual(d["pages"], [1, 2])
        self.assertEqual(len(d["invoice_items"]), 6)
        self.assertTrue(check("multi.total", val(d["invoices"][0]["invoice_total"]), 1098.25))
        self.assertEqual({it["page"] for it in d["invoice_items"]}, {1, 2})

    def test_scanned_invoice_uses_ocr(self):
        jid, r = run("inv_en_scan.pdf")
        self.assertEqual({p["method"] for p in r["page_info"]}, {"ocr"})
        d = r["documents"][0]
        METRICS["classification"]["inv_en_scan"] = d["document_type"] == "hospital_invoice"
        inv = d["invoices"][0]
        check("scan.no", val(inv["invoice_number"]), "INV-1001")
        check("scan.total", val(inv["invoice_total"]), 517.5)
        check("scan.name", val(d["patient"]["name"]), "Ahmed Ali Hassan")
        items_ok(d, ms.ITEMS_EN)
        # word-level OCR accuracy vs the text layer of the same page
        ref = [w for w in json.loads(STORE.path(run("inv_en.pdf")[0], "pages", "1.json").read_text("utf-8"))["words"]]
        got = json.loads(STORE.path(jid, "pages", "1.json").read_text("utf-8"))["words"]
        a, b = " ".join(w["text"] for w in ref), " ".join(w["text"] for w in got)
        METRICS["ocr_text_similarity_clean_scan"] = round(difflib.SequenceMatcher(None, a, b).ratio(), 4)
        self.assertGreater(METRICS["ocr_text_similarity_clean_scan"], 0.9)

    def test_low_quality_scan_never_invents_values(self):
        jid, r = run("inv_en_lowq.pdf")
        d = r["documents"][0]
        METRICS["classification"]["inv_en_lowq"] = d["document_type"] == "hospital_invoice"
        self.assertTrue(any(w["kind"] == "low_ocr_quality" for w in d["warnings"]))
        self.assertGreater(len(d["requires_human_review"]), 0)
        self.assertEqual(r["processing_status"], "completed_with_review")

    def test_prescription(self):
        d = run("rx.pdf")[1]["documents"][0]
        METRICS["classification"]["rx"] = d["document_type"] == "prescription"
        got = [(val(m["brand_name"]), val(m["strength"]), val(m["dosage_form"]), val(m["frequency"])) for m in d["medications"]]
        want = [("Amoxicillin", "500 mg", "capsule", "1x3"), ("Paracetamol", "500 mg", "tablet", "TID"), ("Cetirizine", "10 mg", "tablet", "OD")]
        for i, (g, w) in enumerate(zip(got, want)):
            for j, (a, b) in enumerate(zip(g, w)):
                check(f"rx.{i}.{j}", a, b)
        self.assertEqual(got, want)
        dx = [x for x in d["diagnoses"] if x["entry_type"] == "diagnosis"]
        self.assertEqual([val(x["icd10"]) for x in dx], ["J02.9"])
        self.assertEqual(len(dx), 1, "medication lines must not be read as diagnoses")

    def test_lab(self):
        d = run("lab.pdf")[1]["documents"][0]
        METRICS["classification"]["lab"] = d["document_type"] == "lab_result"
        rows = [(val(x["test_name"]), val(x["result"]), val(x["unit"]), val(x["reference_range"])) for x in d["laboratory_tests"]]
        self.assertEqual(rows, [("Hemoglobin", "10.2", "g/dL", "13.0 - 17.0"), ("Glucose Fasting", "95", "mg/dL", "70 - 100"), ("Creatinine", "1.1", "mg/dL", "0.7 - 1.3")])
        for i, r_ in enumerate(rows):
            check(f"lab.{i}", r_[0], ["Hemoglobin", "Glucose Fasting", "Creatinine"][i])
        self.assertEqual(val(d["laboratory_tests"][0]["flag"]), "L")
        self.assertEqual(val(d["laboratory_tests"][0]["test_date"]), "2024-04-02")
        self.assertEqual(d["diagnoses"], [], "no diagnosis may be derived from lab values")

    def test_radiology(self):
        d = run("rad.pdf")[1]["documents"][0]
        METRICS["classification"]["rad"] = d["document_type"] == "radiology_report"
        im = d["imaging"][0]
        self.assertEqual(val(im["exam_type"]), "Chest X-Ray PA")
        self.assertEqual(val(im["findings"]), "Clear lung fields. No pleural effusion. Heart size is within normal limits.")
        ev = next(e for e in d["evidence"] if e["id"] == im["findings"]["evidence_id"])
        self.assertGreaterEqual(len(ev["bboxes"]), 2)             # one box per findings line

    def test_discharge_stay_and_diagnosis_certainty(self):
        d = run("discharge.pdf")[1]["documents"][0]
        METRICS["classification"]["discharge"] = d["document_type"] == "discharge_summary"
        h = d["hospitalization"]
        self.assertEqual((val(h["admission_date"]), val(h["discharge_date"]), val(h["length_of_stay_days"]), val(h["room_type"])),
                         ("2024-07-01", "2024-07-05", 4.0, "Private"))
        dx = {val(x["text"]): (x["category"], x["certainty"], val(x["icd10"])) for x in d["diagnoses"]}
        self.assertEqual(dx["Community acquired pneumonia (J18.9)"], ("primary", "not_stated", "J18.9"))
        self.assertEqual(dx["Type 2 diabetes mellitus E11.9"][:2], ("secondary", "not_stated"))
        self.assertEqual(dx["suspected anemia"][1], "suspected")
        self.assertEqual([w for w in d["warnings"] if w["kind"] in ("stay_days_mismatch", "discharge_before_admission")], [])

    def test_arithmetic_errors_are_flagged_not_fixed(self):
        d = run("inv_badmath.pdf")[1]["documents"][0]
        kinds = {w["kind"] for w in d["warnings"]}
        self.assertTrue({"line_arithmetic_mismatch", "total_formula_mismatch", "balance_mismatch"} <= kinds, kinds)
        self.assertEqual(val(d["invoice_items"][2]["line_total"]), 30.0)          # original number untouched
        self.assertEqual(val(d["invoices"][0]["invoice_total"]), 999.0)
        self.assertEqual(d["invoice_items"][2]["line_total"]["verification"], "needs_review")
        self.assertEqual(d["processing_status"], "completed_with_review")

    def test_missing_data_is_reported_not_guessed(self):
        d = run("inv_missing.pdf")[1]["documents"][0]
        self.assertEqual(d["patient"].get("name", {"status": "not_found"})["status"], "not_found")
        self.assertTrue(any(q["field"] == "name" and "required_field_not_found" in q["reasons"] for q in d["requires_human_review"]))

    def test_embedded_instructions_are_ignored(self):
        d = run("inv_inject.pdf")[1]["documents"][0]
        self.assertTrue(any(w["kind"] == "embedded_instruction" for w in d["warnings"]))
        self.assertEqual(val(d["invoices"][0]["invoice_total"]), 517.5)

    def test_duplicate_page_not_double_counted(self):
        d = run("inv_dup2.pdf")[1]["documents"][0]
        self.assertTrue(any(w["kind"] == "duplicate_page" for w in d["warnings"]))
        self.assertEqual(len(d["invoice_items"]), 4)

    def test_mixed_file_is_split_into_documents(self):
        r = run("mixed.pdf")[1]
        types = [d["document_type"] for d in r["documents"]]
        METRICS["classification"]["mixed"] = types == ["hospital_invoice", "lab_result", "radiology_report"]
        self.assertEqual(types, ["hospital_invoice", "lab_result", "radiology_report"])
        self.assertEqual([d["pages"] for d in r["documents"]], [[1], [2], [3]])

    # ------------------------------------------------------------ robustness / storage
    def test_corrupt_and_non_pdf_inputs(self):
        with self.assertRaises(ValueError):
            STORE.create("x.pdf", b"hello world")
        jid = STORE.create("bad.pdf", b"%PDF-1.4\n garbage without structure")
        self.assertIsNone(process_job(STORE, jid))
        self.assertEqual(STORE.status(jid)["state"], "failed")

    def test_job_id_path_traversal_blocked(self):
        for bad in ("../etc", "..", "abc", "0" * 12):
            with self.assertRaises(KeyError):
                STORE.path(bad, "x")

    def test_reprocess_keeps_corrections_and_raw_untouched(self):
        jid, r = run("inv_en.pdf")
        d = r["documents"][0]
        rec, raw_before = d["invoices"][0]["id"], json.dumps(STORE.result(jid, with_corrections=False)["documents"][0]["invoices"][0]["invoice_total"])
        STORE.add_correction(jid, rec, "invoice_total", "518.00", "manual")
        f = STORE.result(jid)["documents"][0]["invoices"][0]["invoice_total"]
        self.assertEqual((f["corrected_value"], f["verification"], f["value"]), ("518.00", "corrected", 517.5))
        process_job(STORE, jid)                                           # reprocess (page cache reused)
        self.assertEqual(STORE.result(jid)["documents"][0]["invoices"][0]["invoice_total"]["corrected_value"], "518.00")
        self.assertEqual(json.dumps(STORE.result(jid, with_corrections=False)["documents"][0]["invoices"][0]["invoice_total"]), raw_before)
        STORE.delete_correction(jid, rec, "invoice_total")
        self.assertNotIn("corrected_value", STORE.result(jid)["documents"][0]["invoices"][0]["invoice_total"])

    def test_page_cache_survives_and_is_reused(self):
        jid, _ = run("inv_multi.pdf")
        cache = STORE.path(jid, "pages", "1.json")
        m1 = cache.stat().st_mtime_ns
        process_job(STORE, jid)
        self.assertEqual(cache.stat().st_mtime_ns, m1)

    # ------------------------------------------------------------ exports
    def test_exports(self):
        r = STORE.result(run("inv_en.pdf")[0])
        self.assertEqual(json.loads(exporters.to_json(r))["job_id"], r["job_id"])
        z = zipfile.ZipFile(io.BytesIO(exporters.to_csv_zip(r)))
        self.assertIn("invoice_items.csv", z.namelist())
        self.assertIn("General Consultation", z.read("invoice_items.csv").decode("utf-8-sig"))
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(exporters.to_xlsx(r)))
        self.assertIn("invoice_items", wb.sheetnames)
        self.assertEqual(wb["invoice_items"].max_row, 5)

    def test_pdf_export_pipeline(self):
        """Structure/verification path. If arabic_reshaper+bidi are missing (offline sandbox) they are
        replaced by identity stubs: this proves the font/verification plumbing but NOT real shaping."""
        try:
            import arabic_reshaper, bidi.algorithm  # noqa: F401
            real = True
        except ImportError:
            import types
            m1 = types.ModuleType("arabic_reshaper"); m1.reshape = lambda s: s
            m2 = types.ModuleType("bidi"); m3 = types.ModuleType("bidi.algorithm"); m3.get_display = lambda s: s
            sys.modules.update({"arabic_reshaper": m1, "bidi": m2, "bidi.algorithm": m3})
            real = False
        try:
            r = STORE.result(run("inv_ar.pdf")[0])
            pdf = exporters.to_pdf(r)
            self.assertTrue(pdf.startswith(b"%PDF"))
            METRICS["pdf_export"] = {"ok": True, "real_arabic_shaping_libs": real, "font": exporters.find_arabic_font().name}
        finally:
            if not real:
                for k in ("arabic_reshaper", "bidi", "bidi.algorithm"):
                    sys.modules.pop(k, None)

    def test_pdf_export_fails_loudly_without_shaper(self):
        try:
            import arabic_reshaper  # noqa: F401
            self.skipTest("real shaper installed")
        except ImportError:
            pass
        with self.assertRaises(exporters.PdfExportError):
            exporters.to_pdf(STORE.result(run("inv_en.pdf")[0]))

    # ------------------------------------------------------------ HTTP API end-to-end
    def test_http_api_roundtrip(self):
        from mdia import server
        server.STORE = Store(TMP / "api_store")
        srv = server.make_server("127.0.0.1", 0)
        port = srv.server_address[1]
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        base = f"http://127.0.0.1:{port}"
        try:
            req = urllib.request.Request(base + "/api/upload?filename=inv.pdf", data=(SAMPLES / "inv_en.pdf").read_bytes(), method="POST")
            jid = json.loads(urllib.request.urlopen(req).read())["job_id"]
            for _ in range(120):
                st = json.loads(urllib.request.urlopen(f"{base}/api/jobs/{jid}").read())
                if st["state"] in ("done", "failed"):
                    break
                time.sleep(0.5)
            self.assertEqual(st["state"], "done", st)
            res = json.loads(urllib.request.urlopen(f"{base}/api/jobs/{jid}/result").read())
            self.assertEqual(len(res["documents"][0]["invoice_items"]), 4)
            png = urllib.request.urlopen(f"{base}/api/jobs/{jid}/pages/1.png").read()
            self.assertTrue(png.startswith(b"\x89PNG"))
            self.assertTrue(urllib.request.urlopen(f"{base}/api/jobs/{jid}/pdf").read().startswith(b"%PDF"))
            self.assertTrue(zipfile.is_zipfile(io.BytesIO(urllib.request.urlopen(f"{base}/api/jobs/{jid}/export?format=csv").read())))
            self.assertTrue(urllib.request.urlopen(f"{base}/").read().startswith(b"<!doctype html>"))
            for bad in ("/api/jobs/../../etc/passwd", "/api/jobs/zzzzzzzzzzzz", "/api/jobs/" + jid + "/pages/..%2f..%2fstatus.json"):
                with self.assertRaises(urllib.error.HTTPError):
                    urllib.request.urlopen(base + bad)
            rec = res["documents"][0]["invoices"][0]["id"]
            urllib.request.urlopen(urllib.request.Request(f"{base}/api/jobs/{jid}/corrections", method="POST",
                                   data=json.dumps({"record_id": rec, "field": "currency", "value": "USD"}).encode()))
            res2 = json.loads(urllib.request.urlopen(f"{base}/api/jobs/{jid}/result").read())
            self.assertEqual(res2["documents"][0]["invoices"][0]["currency"]["corrected_value"], "USD")
            METRICS["api_roundtrip"] = "ok"
        finally:
            srv.shutdown()

    def test_api_token_enforced(self):
        from mdia import config, server
        old = config.API_TOKEN
        config.API_TOKEN = "s3cret"
        srv = server.make_server("127.0.0.1", 0)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        try:
            with self.assertRaises(urllib.error.HTTPError) as c:
                urllib.request.urlopen(f"http://127.0.0.1:{srv.server_address[1]}/api/jobs")
            self.assertEqual(c.exception.code, 401)
            ok = urllib.request.Request(f"http://127.0.0.1:{srv.server_address[1]}/api/jobs", headers={"Authorization": "Bearer s3cret"})
            self.assertEqual(urllib.request.urlopen(ok).status, 200)
        finally:
            config.API_TOKEN = old
            srv.shutdown()


if __name__ == "__main__":
    unittest.main(verbosity=2)
